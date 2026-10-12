"""Real PowerPoint with reconciled charts, screenshots and speaker notes."""
import json
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from .common import ROOT,write_json,sha256

def build(root=ROOT):
    f=json.loads((root/'outputs/phase7/validation/verified_facts.json').read_text())
    charts=root/'docs/assets/charts'; charts.mkdir(parents=True,exist_ok=True)
    plt.rcParams.update({'font.size':12,'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white'})
    made=[]
    def save(fig,name):
        path=charts/(name+'.png'); fig.savefig(path,dpi=180,bbox_inches='tight'); plt.close(fig); made.append(path); return path
    frame=pd.read_parquet(root/'outputs/tables/cost_analytics.parquet')
    monthly=frame.groupby('month',as_index=False)[['prescription_items','net_ingredient_cost_gbp']].sum()
    fig,axes=plt.subplots(2,1,figsize=(9,5.7),sharex=True)
    axes[0].plot(monthly.month,monthly.prescription_items/1e6,color='#176B87',lw=2); axes[0].set_ylabel('Items (millions)')
    axes[1].plot(monthly.month,monthly.net_ingredient_cost_gbp/1e6,color='#A85458',lw=2); axes[1].set_ylabel('NIC (£ millions)')
    for ax in axes: ax.grid(alpha=.15)
    fig.autofmt_xdate(); trend=save(fig,'observed_activity_nic')
    d=f['decomposition_2025']; fig,ax=plt.subplots(figsize=(8,4.8)); values=[d['item_volume_effect_gbp'],d['presentation_mix_effect_gbp'],d['within_presentation_cost_per_item_effect_gbp']]
    bars=ax.bar(['Item volume','Presentation mix','Within-presentation\nNIC/item'],[v/1e6 for v in values],color=['#176B87','#479178','#A85458']); ax.axhline(0,color='#718096',lw=1); ax.set_ylabel('Accounting contribution (£ millions)')
    for bar,v in zip(bars,values): ax.annotate(f'{v/1e6:+.2f}',(bar.get_x()+bar.get_width()/2,v/1e6),ha='center',va='bottom' if v>0 else 'top',xytext=(0,4 if v>0 else -4),textcoords='offset points')
    ax.set_ylim(-68,17); decomposition=save(fig,'cost_decomposition')
    records=f['national_holdout_models']; fig,ax=plt.subplots(figsize=(8,4.8)); labels=[r['model'].replace('_',' ').title() for r in records]; vals=[r['MASE'] for r in records]
    bars=ax.barh(labels,vals,color=['#176B87' if r['model']=='damped_ets' else '#9DACBD' for r in records]); ax.set_xlabel('Holdout MASE · common support'); ax.set_xlim(0,1.9)
    for bar,v in zip(bars,vals): ax.text(v+.03,bar.get_y()+bar.get_height()/2,f'{v:.3f}',va='center')
    models=save(fig,'national_model_comparison')
    fig,ax=plt.subplots(figsize=(8,4.8)); baseline=f['nic_2025_gbp']; vals=[baseline,baseline*1.05,baseline*1.05*.95]
    bars=ax.bar(['2025 observed','+5% items','+5% items\n−5% NIC/item'],[v/1e6 for v in vals],color=['#9DACBD','#176B87','#479178']); ax.set_ylabel('NIC (£ millions)'); ax.set_ylim(0,90)
    for b,v in zip(bars,vals): ax.text(b.get_x()+b.get_width()/2,v/1e6+1,f'£{v/1e6:.2f}m',ha='center')
    scenarios=save(fig,'hypothetical_scenarios')
    # The architecture PNG is an actual rasterization of the SVG diagram, not a dashboard mockup.
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='C:/Program Files/Google/Chrome/Application/chrome.exe',headless=True)
        page=browser.new_page(viewport={'width':1200,'height':650}); page.goto((root/'docs/assets/architecture.svg').as_uri()); page.screenshot(path=str(root/'docs/assets/architecture.png')); browser.close()
    architecture=root/'docs/assets/architecture.png'; made.append(architecture)
    slides=[
      ('PharmaLens','Pharmaceutical prescribing intelligence\nIndependent, AI-assisted portfolio project',[f'England · atorvastatin · {f["months"]} observed months','SQL, statistical diagnostics, frozen forecasts and decision support','Public administrative evidence; no proprietary client data'],None,
       'Introduce the project as an independent portfolio implementation supported by public NHSBSA data. State explicitly that Codex assisted substantially with implementation and audit work. Define the business boundary: recorded dispensing items and NHS ingredient cost, not patients or manufacturer revenue. The evidence period is January 2023 through December 2025. Explain the end-to-end connection between source, calculations and dashboard rather than claiming client adoption or savings.'),
      ('01 · Business question','Make the evidence useful without overstating it',['What changed in observed activity and ingredient cost?','Which forecasts and review priorities need caution?','What follows from an explicit accounting assumption?'],None,
       'The questions are descriptive and analytical. Cost trends, reliability review and assumption-driven planning are useful demonstrations, but the project does not identify a market intervention. Explain why an item is not a patient and why NIC is not manufacturer sales. Operational use would require a named decision, authorized business data and user validation.'),
      ('02 · Data scope','Preserve grain, units and provenance',[f'{f["source_records"]:,} source rows → {f["analytical_rows"]:,} monthly rows','7 regions · 17 presentations · 36 observed months','Missing activity stays unknown; physical units remain separate'],None,
       'Describe immutable API snapshots and stable identifiers. The transformed monthly grain aggregates administrative records while preserving item and NIC totals. Distinguish the full region-presentation calendar from observed records. Absence is not converted to zero. Quantity is kept with its source unit because tablets and liquids cannot be pooled as though they were equivalent. Data publication delay and revisions limit any claim of real-time availability.'),
      ('03 · Architecture','A traceable path from source to screen',[],architecture,
       'Walk through source receipts, validation, monthly data, DuckDB facts and views, exploratory diagnostics, origin-constrained forecasts, monitoring scenarios and reporting marts. The Mermaid source and SVG use the same node and edge definition. Cross-cutting validation includes hashes, contracts, independent arithmetic, regression tests and browser checks. The dashboard consumes validated outputs; it does not fit models during a user interaction.'),
      ('04 · Observed findings','Activity rose while NIC fell',['2025 items: 76.86m · +7.58% versus 2024','2025 NIC: £75.97m · −39.93% versus 2024','The measures have different accounting meanings'],trend,
       'Point to the two series and preserve their units. The 2025 item total is 76,860,890 and NIC is £75,973,698.28. The growth comparison is against 2024, not the full-history total. A change in item count does not alone determine ingredient cost because mix and per-item ratios also change. These are observed administrative totals, with no assertion of unique patients or manufacturer revenue.'),
      ('05 · Statistical explanation','An exact identity, with a causal boundary',['Volume: +£7.48m · mix: +£1.25m','Within-presentation NIC/item: −£59.24m','Total observed NIC change: −£50.51m'],decomposition,
       'Explain the three-factor identity: aggregate items times presentation shares times presentation NIC-per-item ratios. Symmetric allocation across factor-change orders handles interactions and reconciles to the observed change. The within-presentation effect outweighs positive volume and mix contributions. This is not evidence that negotiation, expiry or substitution caused the decline. Trend and seasonality diagnostics are descriptive, with only three annual cycles available.'),
      ('06 · Forecasting','Preserve the frozen selection',['Development selection frozen before holdout','National ETS: MASE 0.888 · WAPE 3.684%','Ridge comparator MASE 0.872 is reported honestly'],models,
       'Explain rolling origins and common evaluation support. The national choice was frozen before July–December 2025 holdout targets. ETS improves MASE by about 19.75 percent against evaluated naive on the same support. Ridge has lower holdout MASE, which is disclosed, but replacing ETS after looking at those results would contaminate selection. Horizon counts are six, five and four, and overlapping cells are not independent replications.'),
      ('07 · Reliability limits','Future-dated does not mean already validated',['36 modeled series · 15 require manual review','Intervals unavailable: residual support below stated threshold','Jan–Mar 2026 forecasts use retrospective Dec 2025 cutoff'],None,
       'The forecast scope includes one national, seven regional and twenty-eight region-presentation series. No 2026 actuals validate the operational forecasts. The observation cutoff is not proof of contemporaneous data availability because publication is delayed. Interval calibration requires twenty residuals per horizon under the implemented rule, while the saved support is six, five and four. Independent forecasts can have coherence gaps. These limits are visible rather than hidden by invented intervals.'),
      ('08 · Decision intelligence','Rank analytical review questions',['Five weighted percentile components; policy is illustrative','All 7 regional ranks unstable under saved sensitivity rule','13 presentation priorities unscored; 225 proposed investigations'],None,
       'Describe activity, change, cost, forecast and quality evidence and the explicit weights. Percentiles are relative to their scope and tied values receive documented midranks. Missing required evidence remains unscored. Ranking instability means the order depends on the policy, not that source totals are wrong. The investigation register has source and rule lineage, but the proposed reviews have not been adopted or validated as business actions.'),
      ('09 · Scenario planning','Check assumptions through accounting',['Hypothetical NIC = baseline × volume factor × ratio factor','+5% and −5% produce −0.25%, not break-even','Exact offset for +5% volume: −4.7619% NIC/item'],scenarios,
       'Demonstrate that the simulator is a deterministic accounting map. Using the 2025 observed baseline, five percent more items at fixed NIC per item gives £79,772,383.19. Combining five percent more items and five percent lower ratio gives £75,783,764.03. Equal opposite percentages do not cancel because of their interaction. These are assumptions, not estimated causal effects or realized savings. Mix transfers also require item conservation and nonnegative source quantities.'),
      ('10 · Dashboard and validation','Inspect the evidence, then drill down',['Six pages with scope-aware exports and reliability labels','Independent reconciliation and protected-artifact checks','Real browser and separate runtime-environment evidence'],root/'docs/assets/screenshots/01_executive_overview.png',
       'This is a screenshot of the actual local Streamlit application, not a mockup. Walk through executive, regional, presentation, forecasting, decision and scenario views. Explain what AppTest checks versus Chrome checks and why a matching CSV download matters. The isolated runtime environment tests declared dependencies and packaged marts, while the full original analytical regression is separate. Read final audit counts from the current report rather than inventing a result.'),
      ('11 · Limits and next steps','A defensible portfolio, with clear boundaries',['No manufacturer revenue, causal savings or clinical recommendations','More history, release-aware evaluation and user-reviewed policies','GitHub/cloud publication and code licence remain owner decisions'],None,
       'Close by naming limitations before future additions. Additional history and prospective data could strengthen evaluation and calibration. Authorized commercial fields would support a separate, correctly defined business analysis. Domain users should validate priority policies and investigation usefulness. No remote repository, cloud release or actual GitHub Actions execution is claimed. Attribution to NHSBSA and the Open Government Licence is supplied; source-code licence choice remains an owner decision.'),
    ]
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    navy=RGBColor.from_string('142B45'); teal=RGBColor.from_string('176B87'); grey=RGBColor.from_string('52647A')
    def text(slide,x,y,w,h,content,size,color=navy,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True
        for i,line in enumerate(content.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(size); p.font.color.rgb=color; p.font.bold=bold; p.space_after=Pt(14)
        return box
    for index,(label,title,bullets,picture,notes) in enumerate(slides,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        text(slide,.55,.3,12.1,.4,label,15,teal,True); text(slide,.55,.88,12.1,1,title,30,navy,True)
        if picture:
            left=4.9 if bullets else .55; available=7.85 if bullets else 12.2
            with Image.open(picture) as im: ratio=im.width/im.height
            w=min(available,4.9*ratio); h=w/ratio
            slide.shapes.add_picture(str(picture),Inches(left+(available-w)/2),Inches(2.05),width=Inches(w),height=Inches(h))
        if bullets: text(slide,.65,2.05,3.95 if picture else 11.8,4.55,'\n'.join('• '+b for b in bullets),20)
        text(slide,.55,7.08,11.8,.25,'PharmaLens · public NHSBSA data · independent AI-assisted portfolio · October 2026',10,grey)
        text(slide,12.3,7.08,.5,.25,f'{index:02d}',10,grey)
        slide.notes_slide.notes_text_frame.text=notes
    dest=root/'deliverables/Phase7_Axtria_Project_Presentation.pptx'; dest.parent.mkdir(exist_ok=True); prs.save(dest)
    write_json(root/'outputs/phase7/validation/presentation_inventory.json',{'slides':len(slides),'speaker_notes':len(slides),'source':'verified_facts.json and immutable cost_analytics.parquet','chart_hashes':{p.relative_to(root).as_posix():sha256(p) for p in made},'pptx_sha256':sha256(dest),'native_powerpoint_render_review':False})
    print(f'Created {dest.name}: {len(slides)} slides and speaker notes')

if __name__=='__main__': build()
